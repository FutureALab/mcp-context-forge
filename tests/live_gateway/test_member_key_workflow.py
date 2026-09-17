"""Exercise member key issuance and analytics against a live isolated gateway.

Set MEMBER_TEST_ADMIN_TOKEN, MEMBER_TEST_EMAIL, MEMBER_TEST_PASSWORD, MEMBER_TEST_TEAM, and MEMBER_TEST_SERVER.
Set MEMBER_TEST_IMPORT_EMAIL to test an existing account outside the selected team.
Copyright contributors to the MCP-CONTEXT-FORGE project.
SPDX-License-Identifier: Apache-2.0
"""

# Standard
import io
from datetime import datetime, timedelta
import os
import time
import unittest

# Third-Party
import httpx
import openpyxl


class TestMemberKeyWorkflow(unittest.TestCase):
    """Validate public membership checks, MCP use, imports, and aggregation."""

    def test_full_workflow(self):
        """Run the public and administrative flows with real HTTP requests."""
        token = os.getenv("MEMBER_TEST_ADMIN_TOKEN")
        if not token:
            self.skipTest("Set MEMBER_TEST_ADMIN_TOKEN for an isolated gateway")
        origin = os.getenv("GATEWAY_URL", "http://127.0.0.1:4445")
        email, team, server = (os.environ[k] for k in ("MEMBER_TEST_EMAIL", "MEMBER_TEST_TEAM", "MEMBER_TEST_SERVER"))
        password = os.environ["MEMBER_TEST_PASSWORD"]
        with httpx.Client(base_url=origin, trust_env=False, timeout=30) as public, httpx.Client(base_url=origin, trust_env=False, timeout=30) as admin:
            page = public.get("/admin/api-key")
            self.assertEqual(page.status_code, 200)
            self.assertIn('id="member-key-form"', page.text)
            self.assertIn("no-store", page.headers["cache-control"])
            self.assertEqual(public.post("/admin/api-key/issue", json={"password": password, "email": email, "team_id": team, "server_id": server}).status_code, 403)
            public.headers["X-CSRF-Token"] = public.cookies["mcpgateway_csrf_token"]
            public.headers["Origin"] = origin
            options = public.post("/admin/api-key/options", json={"password": password, "email": email})
            self.assertEqual(options.status_code, 200, options.text)
            self.assertIn(server, [s["id"] for s in options.json()["servers"]])
            public_server = os.getenv("MEMBER_TEST_PUBLIC_SERVER")
            if public_server:
                public_option = next(s for s in options.json()["servers"] if s["id"] == public_server)
                self.assertIn(team, public_option["eligible_team_ids"])
            for endpoint in ("options", "issue"):
                wrong = public.post(f"/admin/api-key/{endpoint}", json={"email": email, "password": "incorrect", "team_id": team, "server_id": server})
                self.assertEqual(wrong.status_code, 403)
                missing = public.post(f"/admin/api-key/{endpoint}", json={"email": email, "team_id": team, "server_id": server})
                self.assertEqual(missing.status_code, 422)
            bad = public.post("/admin/api-key/issue", json={"password": password, "email": email, "team_id": "wrong-team", "server_id": server})
            self.assertEqual(bad.status_code, 403, bad.text)
            issued = public.post("/admin/api-key/issue", json={"password": password, "email": email, "team_id": team, "server_id": server, "days": 7})
            self.assertEqual(issued.status_code, 200, issued.text if issued.status_code != 200 else "")
            key = issued.json()["api_key"]
            limit = datetime.fromisoformat(issued.json()["created_at"]) + timedelta(days=365)
            for days in (180, 365, 30):
                renewed = public.post("/admin/api-key/issue", json={"password": password, "email": email, "team_id": team, "server_id": server, "days": days})
                self.assertEqual(renewed.status_code, 200)
                self.assertTrue(renewed.json()["renewed"])
                self.assertEqual(renewed.json()["token_id"], issued.json()["token_id"])
                self.assertEqual(renewed.json()["api_key"], "")
                self.assertEqual(datetime.fromisoformat(renewed.json()["expires_at"]), min(datetime.fromisoformat(issued.json()["expires_at"]) + timedelta(days=days), limit))
                issued = renewed
            self.assertIn("no-store", issued.headers["cache-control"])
            self.assertNotIn("jwt_token", public.cookies)
            mcp_headers = {"Authorization": f"Bearer {key}", "Accept": "application/json, text/event-stream"}
            response = public.post(
                f"/servers/{server}/mcp",
                headers=mcp_headers,
                json={"jsonrpc": "2.0", "id": 1, "method": "initialize", "params": {"protocolVersion": "2025-03-26", "capabilities": {}, "clientInfo": {"name": "member-test", "version": "1"}}},
            )
            self.assertEqual(response.status_code, 200, response.text)
            if "mcp-session-id" in response.headers:
                mcp_headers["Mcp-Session-Id"] = response.headers["mcp-session-id"]
            public.post(f"/servers/{server}/mcp", headers=mcp_headers, json={"jsonrpc": "2.0", "method": "notifications/initialized"})
            response = public.post(f"/servers/{server}/mcp", headers=mcp_headers, json={"jsonrpc": "2.0", "id": 2, "method": "tools/list"})
            self.assertEqual(response.status_code, 200, response.text)
            self.assertIn("tools", response.text)
            tool = os.getenv("MEMBER_TEST_TOOL")
            if tool:
                response = public.post(f"/servers/{server}/mcp/", headers=mcp_headers, json={"jsonrpc": "2.0", "id": 3, "method": "tools/call", "params": {"name": tool, "arguments": {}}})
                self.assertEqual(response.status_code, 200, response.text)
                self.assertIn("pong", response.text)
            admin.headers["Authorization"] = f"Bearer {token}"
            panel = admin.get("/admin/member-usage")
            self.assertEqual(panel.status_code, 303, panel.text[:200])
            self.assertTrue(panel.headers["location"].endswith("#member-usage"))
            panel = admin.get("/admin/")
            self.assertEqual(panel.status_code, 200, panel.text[:200])
            self.assertIn('id="member-usage-panel"', panel.text)
            self.assertIn('id="server-members-dialog"', panel.text)
            self.assertIn('id="tab-member-usage"', panel.text)
            admin.headers["Origin"] = origin
            admin.headers["X-CSRF-Token"] = admin.cookies["mcpgateway_csrf_token"]
            reveal_path = f"/tokens/{issued.json()['token_id']}/reveal"
            revealed = admin.post(reveal_path)
            self.assertEqual(revealed.status_code, 200)
            self.assertEqual(revealed.json()["access_token"], key)
            self.assertIn("no-store", revealed.headers["cache-control"])
            token_list = admin.get("/tokens").text
            self.assertNotIn(key, token_list)
            self.assertNotIn("encrypted_token", token_list)
            self.assertIn(public.post(reveal_path).status_code, (401, 403))
            self.assertIn(public.post(reveal_path, headers={"Authorization": f"Bearer {key}"}).status_code, (401, 403))
            if public_server and os.getenv("MEMBER_TEST_PERSONAL_PUBLIC"):
                choices = admin.get("/admin/member-usage/servers").json()
                selected = next(item for item in choices["servers"] if item["id"] == public_server)
                self.assertTrue(selected["is_personal"])
                self.assertIn(team, [item["id"] for item in choices["teams"]])
                workbook = openpyxl.Workbook()
                workbook.active.append(["邮箱"])
                workbook.active.append(["outsider@example.com"])
                content = io.BytesIO()
                workbook.save(content)
                workbook.close()
                upload = {"file": ("members.xlsx", content.getvalue())}
                rejected = admin.post(f"/admin/member-usage/import/{public_server}", files=upload)
                self.assertEqual(rejected.status_code, 400)
                accepted = admin.post(f"/admin/member-usage/import/{public_server}", files=upload, data={"team_id": team})
                self.assertEqual(accepted.status_code, 200)
                self.assertEqual(accepted.json()["results"][0]["status"], "added")
                bulk = admin.post("/admin/member-usage/keys", json={"server_id": public_server, "team_id": team, "days": 180})
                self.assertEqual(bulk.status_code, 200)
                self.assertIn("outsider@example.com", [row["email"] for row in bulk.json()["results"] if row["status"] == "created"])
            token_request = {"name": "live-renewable-key", "team_id": team, "scope": {"server_id": server}, "expires_in_days": 180}
            created_token = admin.post("/tokens", json=token_request)
            self.assertEqual(created_token.status_code, 201)
            self.assertTrue(created_token.json()["access_token"])
            token_request["expires_in_days"] = 365
            renewed_token = admin.post(f"/tokens/teams/{team}", json=token_request)
            self.assertEqual(renewed_token.status_code, 201)
            self.assertTrue(renewed_token.json()["renewed"])
            self.assertEqual(renewed_token.json()["access_token"], "")
            self.assertEqual(created_token.json()["token"]["id"], renewed_token.json()["token"]["id"])
            template = admin.get("/admin/users/import-template")
            self.assertEqual(template.status_code, 200)
            workbook = openpyxl.load_workbook(io.BytesIO(template.content))
            self.assertEqual(workbook.active.cell(1, 1).value, "邮箱")
            workbook.active.append(["bulk-import@example.com", "导入测试用户", None, "否", team])
            content = io.BytesIO()
            workbook.save(content)
            workbook.close()
            for _ in range(2):
                imported_users = admin.post("/admin/users/bulk-import", files={"file": ("users.xlsx", content.getvalue())})
                self.assertEqual(imported_users.status_code, 200)
                self.assertIn("失败 0", imported_users.text)
            self.assertIn("账号已存在", imported_users.text)
            imported = os.getenv("MEMBER_TEST_IMPORT_EMAIL")
            if imported:
                workbook = openpyxl.Workbook()
                workbook.active.append(["邮箱"])
                for account in (email, imported, "missing-import@example.com"):
                    workbook.active.append([account])
                content = io.BytesIO()
                workbook.save(content)
                workbook.close()
                response = admin.post(f"/admin/member-usage/import/{server}", files={"file": ("members.xlsx", content.getvalue())})
                self.assertEqual(response.status_code, 200, response.text)
                outcomes = {r["email"]: r["status"] for r in response.json()["results"]}
                self.assertEqual(outcomes[email], "skipped")
                self.assertIn(outcomes[imported], ("added", "skipped"))
                self.assertEqual(outcomes["missing-import@example.com"], "failed")
            response = admin.post("/admin/member-usage/keys", json={"server_id": server, "days": 7})
            self.assertEqual(response.status_code, 200, response.text if response.status_code != 200 else "")
            keys = [r["api_key"] for r in response.json()["results"] if r["status"] == "created"]
            self.assertGreaterEqual(len(keys), 1)
            self.assertEqual(len(keys), len(set(keys)))
            exported = admin.post("/admin/member-usage/keys/export", json={"results": response.json()["results"]})
            self.assertEqual(exported.status_code, 200)
            self.assertIn(".xlsx", exported.headers["content-disposition"])
            self.assertIn("no-store", exported.headers["cache-control"])
            workbook = openpyxl.load_workbook(io.BytesIO(exported.content))
            self.assertEqual(workbook.active.cell(1, 6).value, "API Key")
            self.assertGreaterEqual(workbook.active.max_row, 2)
            workbook.close()
            for _ in range(10):
                response = admin.get("/admin/member-usage/data", params={"server_id": server, "email": email})
                self.assertEqual(response.status_code, 200, response.text[:200])
                data = response.json()
                if any(m["method"] == "tools/list" for m in data["methods"]):
                    break
                time.sleep(0.2)
            self.assertTrue(any(m["method"] == "tools/list" for m in data["methods"]), data["methods"])
            self.assertGreaterEqual(data["summary"]["calls"], 2)
            if tool:
                call = next(r for r in data["recent"] if r["resource"] == tool)
                self.assertEqual(call["method"], "tools/call")
                self.assertEqual(call["outcome"], "success")
                self.assertIsNotNone(call["timestamp"])
                filtered = admin.get("/admin/member-usage/data", params={"server_id": server, "method": tool}).json()
                self.assertTrue(filtered["recent"])
                self.assertTrue(all(r["resource"] == tool for r in filtered["recent"]))
            self.assertIsNone(data["summary"]["input_tokens"])
            self.assertNotIn(key, response.text)
            denied = public.get("/admin/member-usage/data", headers={"Authorization": f"Bearer {key}"})
            self.assertIn(denied.status_code, (401, 403))


if __name__ == "__main__":
    unittest.main()
