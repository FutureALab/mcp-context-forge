"""Check membership boundaries, token scoping, and MCP usage extraction.

Copyright contributors to the MCP-CONTEXT-FORGE project.
SPDX-License-Identifier: Apache-2.0
"""

# Standard
import importlib
import io
import json
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock, patch

# Third-Party
from fastapi import HTTPException
import jwt
import openpyxl
from sqlalchemy import create_engine, inspect, text
from sqlalchemy.orm import Session

# First-Party
from mcpgateway.db import Base, EmailTeam, EmailTeamMember, EmailUser, Role, Server, TokenUsageLog, UserRole
from mcpgateway.services.mcp_usage_capture import MCPUsageCapture
from mcpgateway.services.member_usage_service import member_usage_summary, summarize
from mcpgateway.services.server_member_service import issue_member_key, member_server


class TestMemberKeys(unittest.IsolatedAsyncioTestCase):
    """Use real persistence and JWT issuance for membership regression cases."""

    def setUp(self):
        """Create two isolated teams and one server."""
        self.engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(self.engine)
        self.db = Session(self.engine)
        self.db.add_all([EmailUser(email="member@example.com"), EmailUser(email="other@example.com")])
        self.db.flush()
        self.db.add_all([EmailTeam(id="team-a", name="A", slug="team-a", created_by="member@example.com"), EmailTeam(id="team-b", name="B", slug="team-b", created_by="other@example.com")])
        self.db.flush()
        self.db.add(EmailTeamMember(team_id="team-a", user_email="member@example.com"))
        self.db.add(Server(id="server-a", name="A", team_id="team-a", visibility="team", owner_email="member@example.com"))
        self.db.add(Role(id="role-a", name="member-test", scope="team", created_by="member@example.com", permissions=["servers.use", "tools.read", "tools.execute"]))
        self.db.flush()
        self.db.add(UserRole(user_email="member@example.com", role_id="role-a", scope="team", scope_id="team-a", granted_by="member@example.com"))
        self.db.commit()

    def tearDown(self):
        """Release the isolated database."""
        self.db.close()
        self.engine.dispose()

    async def test_reuse_scoped_keys(self):
        """Renew matching JWTs without granting management permissions."""
        with patch("mcpgateway.services.server_member_service.get_audit_trail_service"):
            first = await issue_member_key(self.db, "MEMBER@example.com", "team-a", "server-a", 7, "self-service")
            second = await issue_member_key(self.db, "member@example.com", "team-a", "server-a", 7, "self-service")
        self.assertEqual(first["token_id"], second["token_id"])
        self.assertTrue(second["renewed"])
        self.assertEqual(second["api_key"], "")
        self.assertGreater(second["expires_at"], first["expires_at"])
        claims = jwt.decode(first["api_key"], options={"verify_signature": False})
        self.assertEqual(claims["teams"], ["team-a"])
        self.assertEqual(claims["scopes"]["server_id"], "server-a")
        self.assertNotIn("tokens.create", claims["scopes"]["permissions"])

    async def test_public_key_without_team_and_private_resource_denials(self):
        """Allow a registered viewer to use public MCPs without a team."""
        from mcpgateway.routers.server_members import member_options_for_user

        self.db.add(EmailUser(email="public@example.com"))
        self.db.add(Role(id="role-public", name="public-viewer-test", scope="global", created_by="member@example.com", permissions=["servers.use", "tools.read"]))
        self.db.flush()
        self.db.add(UserRole(user_email="public@example.com", role_id="role-public", scope="global", granted_by="member@example.com"))
        self.db.add(Server(id="server-public", name="Public", team_id="team-a", visibility="public", owner_email="member@example.com"))
        self.db.add(Server(id="server-private", name="Private", team_id="team-a", visibility="private", owner_email="member@example.com"))
        self.db.commit()

        user = self.db.query(EmailUser).filter_by(email="public@example.com").one()
        options = json.loads((await member_options_for_user(self.db, user)).body)
        self.assertEqual(options["teams"], [])
        self.assertEqual([server["id"] for server in options["servers"]], ["server-public"])
        self.assertTrue(options["servers"][0]["public_access"])
        self.assertEqual(options["servers"][0]["eligible_team_ids"], [])

        with patch("mcpgateway.services.server_member_service.get_audit_trail_service"):
            issued = await issue_member_key(self.db, user.email, None, "server-public", 7, "self-service")
        claims = jwt.decode(issued["api_key"], options={"verify_signature": False})
        self.assertIsNone(issued["team_id"])
        self.assertIsNone(claims["teams"])
        self.assertFalse(claims["user"]["is_admin"])
        self.assertEqual(claims["scopes"]["server_id"], "server-public")
        self.assertEqual(claims["scopes"]["permissions"], ["servers.use", "tools.read"])

        for server_id in ("server-a", "server-private"):
            with self.subTest(server_id=server_id), self.assertRaises(HTTPException) as error:
                await issue_member_key(self.db, user.email, None, server_id, 7, "self-service")
            self.assertEqual(error.exception.status_code, 403)
        with self.assertRaises(HTTPException):
            await issue_member_key(self.db, "missing@example.com", None, "server-public", 7, "self-service")

        self.db.add(EmailUser(email="unprivileged@example.com"))
        self.db.commit()
        with self.assertRaises(HTTPException):
            await issue_member_key(self.db, "unprivileged@example.com", None, "server-public", 7, "self-service")

        public_server = self.db.get(Server, "server-public")
        public_server.enabled = False
        self.db.commit()
        with self.assertRaises(HTTPException):
            await issue_member_key(self.db, user.email, None, "server-public", 7, "self-service")
        public_server.enabled = True
        self.db.commit()

        user.is_admin = True
        self.db.commit()
        with self.assertRaises(HTTPException):
            await issue_member_key(self.db, user.email, None, "server-public", 7, "self-service")

    async def test_personal_team_remains_available_without_shared_team(self):
        """List and issue keys for a member's personal workspace."""
        from mcpgateway.routers.server_members import member_options_for_user

        team = self.db.get(EmailTeam, "team-a")
        team.is_personal = True
        self.db.commit()
        user = self.db.query(EmailUser).filter_by(email="member@example.com").one()
        options = json.loads((await member_options_for_user(self.db, user)).body)
        self.assertEqual(options["teams"], [{"id": "team-a", "name": "A", "is_personal": True}])
        self.assertEqual(options["servers"][0]["eligible_team_ids"], ["team-a"])
        with patch("mcpgateway.services.server_member_service.get_audit_trail_service"):
            issued = await issue_member_key(self.db, user.email, "team-a", "server-a", 7, "self-service")
        self.assertEqual(issued["team_id"], "team-a")

    async def test_server_key_inventory_includes_all_owners_and_exports_existing_keys(self):
        """Keep historical and cross-team records while exporting exact stored material."""
        from mcpgateway.db import EmailApiToken
        from mcpgateway.routers.server_members import export_member_key_history, member_key_detail, member_key_history

        admin = {"email": "admin@example.com", "auth_method": "jwt", "is_admin": True, "token_teams": None}
        with patch("mcpgateway.services.server_member_service.get_audit_trail_service"):
            issued = await issue_member_key(self.db, "member@example.com", "team-a", "server-a", 7, "test")
        record = self.db.get(EmailApiToken, issued["token_id"])
        original_expiry = record.expires_at
        self.db.add(EmailApiToken(id="legacy", user_email="other@example.com", team_id="team-b", name="=1+1", token_hash="historical", server_id="server-a", is_active=False))
        self.db.add(EmailApiToken(id="unscoped", user_email="other@example.com", name="general", token_hash="general"))
        self.db.commit()
        response = await member_key_history.__wrapped__("server-a", 0, 1, admin, self.db)
        payload = json.loads(response.body)
        self.assertEqual(payload["total"], 2)
        self.assertEqual(len(payload["results"]), 1)
        self.assertNotIn(issued["api_key"], response.body.decode())
        self.assertNotIn("encrypted_token", response.body.decode())
        with patch("mcpgateway.routers.server_members.get_audit_trail_service") as audit:
            detail = await member_key_detail.__wrapped__("server-a", issued["token_id"], admin, self.db)
            exported = await export_member_key_history.__wrapped__("server-a", admin, self.db)
        self.assertEqual(json.loads(detail.body)["api_key"], issued["api_key"])
        self.assertNotIn(issued["api_key"], str(audit.mock_calls))
        workbook = openpyxl.load_workbook(io.BytesIO(exported.body))
        rows = list(workbook.active.iter_rows())
        headers = [cell.value for cell in rows[0]]
        records = [{headers[i]: cell.value for i, cell in enumerate(row)} for row in rows[1:]]
        self.assertEqual({row["成员邮箱"] for row in records}, {"member@example.com", "other@example.com"})
        self.assertIn(issued["api_key"], [row["API Key"] for row in records])
        legacy = next(row for row in records if row["Key ID"] == "legacy")
        self.assertEqual(legacy["状态"], "已停用")
        self.assertIn("无法恢复", legacy["原文情况"])
        name_cell = next(row[headers.index("Key 名称")] for row in rows[1:] if row[headers.index("Key ID")].value == "legacy")
        self.assertEqual(name_cell.data_type, "s")
        workbook.close()
        self.assertEqual(exported.headers["cache-control"], "no-store")
        self.assertEqual(record.expires_at, original_expiry)
        from datetime import timedelta
        from mcpgateway.db import TokenRevocation, utc_now
        from mcpgateway.services.server_key_inventory import inventory_query, inventory_record

        record.expires_at = utc_now() - timedelta(days=1)
        self.db.commit()
        query = inventory_query("server-a").where(EmailApiToken.id == record.id)
        self.assertEqual(inventory_record(self.db.execute(query).first())["status"], "已到期")
        self.db.add(TokenRevocation(jti=record.jti, revoked_by="member@example.com", reason="test revocation"))
        self.db.commit()
        revoked = inventory_record(self.db.execute(query).first())
        self.assertEqual(revoked["status"], "已吊销")
        self.assertEqual(revoked["revocation_reason"], "test revocation")

    async def test_server_key_inventory_rejects_non_platform_sessions_and_wrong_server(self):
        """Deny nonadministrative sessions and mismatched server/key pairs."""
        from mcpgateway.routers.server_members import export_member_key_history, member_key_detail, member_key_history

        with patch("mcpgateway.services.server_member_service.get_audit_trail_service"):
            issued = await issue_member_key(self.db, "member@example.com", "team-a", "server-a", 7, "test")
        for user in [
            {"email": "member@example.com"},
            {"email": "member@example.com", "auth_method": "jwt", "is_admin": False},
            {"email": "member@example.com", "auth_method": "api_token", "is_admin": True},
            {"email": "member@example.com", "auth_method": "jwt", "is_admin": True, "token_teams": ["team-a"]},
        ]:
            for function, args in [(member_key_history, ("server-a", 0, 50)), (export_member_key_history, ("server-a",)), (member_key_detail, ("server-a", issued["token_id"]))]:
                with self.subTest(user=user, endpoint=function.__name__), self.assertRaises(HTTPException):
                    await function.__wrapped__(*args, user, self.db)
        admin = {"email": "admin@example.com", "auth_method": "jwt", "is_admin": True}
        with self.assertRaises(HTTPException) as error:
            await member_key_detail.__wrapped__("wrong-server", issued["token_id"], admin, self.db)
        self.assertEqual(error.exception.status_code, 404)

    async def test_reveal_persists_encrypted_key_across_sessions_and_renewal(self):
        """Reveal the original key after reopening storage and renewing its registry expiry."""
        from mcpgateway.db import EmailApiToken
        from mcpgateway.routers.tokens import reveal_token

        with patch("mcpgateway.services.server_member_service.get_audit_trail_service"):
            issued = await issue_member_key(self.db, "member@example.com", "team-a", "server-a", 7, "test")
            await issue_member_key(self.db, "member@example.com", "team-a", "server-a", 7, "test")
        stored = self.db.get(EmailApiToken, issued["token_id"])
        self.assertTrue(stored.encrypted_token.startswith("v2:"))
        self.assertNotIn(issued["api_key"], stored.encrypted_token)
        self.db.close()
        self.db = Session(self.engine)
        with patch("mcpgateway.routers.tokens.get_audit_trail_service") as audit:
            response = await reveal_token.__wrapped__(issued["token_id"], {"email": "member@example.com", "auth_method": "jwt", "token_teams": ["team-a"]}, self.db)
        self.assertEqual(json.loads(response.body)["access_token"], issued["api_key"])
        self.assertEqual(response.headers["cache-control"], "no-store")
        self.assertNotIn(issued["api_key"], str(audit.mock_calls))

    async def test_reveal_rejects_other_owners_and_narrowed_sessions(self):
        """Deny unauthenticated, API-token, wrong-owner, and wrong-team retrieval."""
        from mcpgateway.routers.tokens import reveal_token

        with patch("mcpgateway.services.server_member_service.get_audit_trail_service"):
            issued = await issue_member_key(self.db, "member@example.com", "team-a", "server-a", 7, "test")
        for user in [
            {"email": "member@example.com"},
            {"email": "member@example.com", "auth_method": "anonymous"},
            {"email": "member@example.com", "auth_method": "api_token"},
            {"email": "other@example.com", "auth_method": "jwt"},
            {"email": "member@example.com", "auth_method": "jwt", "token_teams": []},
            {"email": "other@example.com", "auth_method": "jwt", "is_admin": True, "token_teams": ["team-b"]},
        ]:
            with self.subTest(user=user), self.assertRaises(HTTPException):
                await reveal_token.__wrapped__(issued["token_id"], user, self.db)
        with patch("mcpgateway.routers.tokens.get_audit_trail_service"):
            response = await reveal_token.__wrapped__(issued["token_id"], {"email": "other@example.com", "auth_method": "jwt", "is_admin": True, "token_teams": None}, self.db)
        self.assertEqual(json.loads(response.body)["access_token"], issued["api_key"])

    async def test_historical_key_is_unrecoverable_without_replacing_it(self):
        """Return an explicit historical notice without minting a replacement."""
        from mcpgateway.db import EmailApiToken
        from mcpgateway.routers.tokens import reveal_token

        with patch("mcpgateway.services.server_member_service.get_audit_trail_service"):
            issued = await issue_member_key(self.db, "member@example.com", "team-a", "server-a", 7, "test")
        record = self.db.get(EmailApiToken, issued["token_id"])
        record.encrypted_token = None
        original_hash = record.token_hash
        self.db.commit()
        response = await reveal_token.__wrapped__(issued["token_id"], {"email": "member@example.com", "auth_method": "jwt"}, self.db)
        self.assertIsNone(json.loads(response.body)["access_token"])
        self.assertIn("无法恢复", json.loads(response.body)["message"])
        self.assertEqual(record.token_hash, original_hash)

    async def test_expired_and_disabled_keys_create_replacements(self):
        """Create a replacement when the previous key cannot be renewed."""
        # Standard
        from datetime import timedelta

        # First-Party
        from mcpgateway.db import EmailApiToken, utc_now

        with patch("mcpgateway.services.server_member_service.get_audit_trail_service"):
            first = await issue_member_key(self.db, "member@example.com", "team-a", "server-a", 180, "self-service")
            self.db.get(EmailApiToken, first["token_id"]).expires_at = utc_now() - timedelta(seconds=1)
            self.db.commit()
            second = await issue_member_key(self.db, "member@example.com", "team-a", "server-a", 365, "self-service")
            self.assertNotEqual(first["token_id"], second["token_id"])
            self.assertTrue(second["api_key"])
            self.db.get(EmailApiToken, second["token_id"]).is_active = False
            self.db.commit()
            third = await issue_member_key(self.db, "member@example.com", "team-a", "server-a", 180, "self-service")
            self.assertNotEqual(second["token_id"], third["token_id"])
            self.assertFalse(third["renewed"])

    async def test_lifetime_is_capped_without_replacing_key(self):
        """Clamp renewal to creation plus 365 days and preserve existing longer keys."""
        # Standard
        from datetime import datetime, timedelta

        # First-Party
        from mcpgateway.db import EmailApiToken

        with patch("mcpgateway.services.server_member_service.get_audit_trail_service"):
            first = await issue_member_key(self.db, "member@example.com", "team-a", "server-a", 180, "self-service")
            limit = datetime.fromisoformat(first["created_at"]) + timedelta(days=365)
            for days in (365, 180):
                renewed = await issue_member_key(self.db, "member@example.com", "team-a", "server-a", days, "self-service")
                self.assertEqual(renewed["token_id"], first["token_id"])
                self.assertEqual(datetime.fromisoformat(renewed["expires_at"]), limit)
                self.assertTrue(renewed["limit_reached"])
                self.assertEqual(renewed["api_key"], "")
            record = self.db.get(EmailApiToken, first["token_id"])
            record.expires_at = limit + timedelta(days=30)
            self.db.commit()
            unchanged = await issue_member_key(self.db, "member@example.com", "team-a", "server-a", 365, "self-service")
            self.assertEqual(datetime.fromisoformat(unchanged["expires_at"]), limit + timedelta(days=30))

    def test_management_team_preserves_visibility_boundaries(self):
        """Allow shared-team selection for public servers and reject unsafe targets."""
        # First-Party
        from mcpgateway.routers.server_members import management_team

        server = self.db.get(Server, "server-a")
        with self.assertRaises(HTTPException):
            management_team(self.db, server, "team-b")
        server.visibility = "public"
        self.assertEqual(management_team(self.db, server, "team-b").id, "team-b")
        for target in ("missing", "team-a"):
            self.db.get(EmailTeam, "team-a").is_personal = True
            with self.assertRaises(HTTPException):
                management_team(self.db, server, target)
        self.db.get(EmailTeam, "team-b").is_active = False
        with self.assertRaises(HTTPException):
            management_team(self.db, server, "team-b")

    async def test_export_is_excel_and_formula_safe(self):
        """Export text cells and reject nonadministrative callers."""
        # Standard
        import io

        # Third-Party
        import openpyxl

        # First-Party
        from mcpgateway.routers.server_members import KeyExportRequest, member_keys_export

        body = KeyExportRequest(results=[{"email": "member@example.com", "server_name": "=1+1", "api_key": "test-only-key"}])
        response = await member_keys_export.__wrapped__(body=body, user={"email": "admin@example.com", "is_admin": True, "auth_method": "session", "token_teams": None})
        workbook = openpyxl.load_workbook(io.BytesIO(response.body))
        self.assertEqual(workbook.active.cell(2, 2).data_type, "s")
        self.assertEqual(workbook.active.cell(2, 6).value, "test-only-key")
        self.assertEqual(response.headers["cache-control"], "no-store")
        workbook.close()
        with self.assertRaises(HTTPException):
            await member_keys_export.__wrapped__(body=body, user={"email": "member@example.com", "is_admin": False, "auth_method": "session"})

    async def test_renewed_key_after_signed_expiration(self):
        """Accept registry-renewed keys while rejecting expired, revoked, and forged keys."""
        # Standard
        from datetime import timedelta
        import hashlib
        import time

        # First-Party
        from mcpgateway.db import EmailApiToken, TokenRevocation, utc_now
        from mcpgateway.utils.jwt_config_helper import get_jwt_public_key_or_secret
        from mcpgateway.utils.verify_credentials import verify_jwt_token

        with patch("mcpgateway.services.server_member_service.get_audit_trail_service"):
            issued = await issue_member_key(self.db, "member@example.com", "team-a", "server-a", 180, "self-service")
        claims = jwt.decode(issued["api_key"], options={"verify_signature": False})
        claims["exp"] = int(time.time()) - 60
        raw = jwt.encode(claims, get_jwt_public_key_or_secret(), algorithm="HS256")
        record = self.db.get(EmailApiToken, issued["token_id"])
        record.token_hash = hashlib.sha256(raw.encode()).hexdigest()
        self.db.commit()

        async def inline(func, *args):
            return func(*args)

        with patch("mcpgateway.db.SessionLocal", side_effect=lambda: Session(self.engine)), patch("mcpgateway.utils.verify_credentials.asyncio.to_thread", side_effect=inline):
            verified = await verify_jwt_token(raw)
            self.assertGreater(verified["exp"], time.time())
            record.is_active = False
            self.db.commit()
            with self.assertRaises(HTTPException):
                await verify_jwt_token(raw)
            record.is_active = True
            self.db.commit()
            with self.assertRaises(HTTPException):
                await verify_jwt_token(raw + "tampered")
            record.expires_at = utc_now() - timedelta(seconds=1)
            self.db.commit()
            with self.assertRaises(HTTPException):
                await verify_jwt_token(raw)
            record.expires_at = utc_now() + timedelta(days=365)
            self.db.add(TokenRevocation(jti=record.jti, revoked_by="member@example.com", reason="test"))
            self.db.commit()
            with self.assertRaises(HTTPException):
                await verify_jwt_token(raw)
            session_claims = {**claims, "token_use": "session"}
            session_raw = jwt.encode(session_claims, get_jwt_public_key_or_secret(), algorithm="HS256")
            with self.assertRaises(HTTPException):
                await verify_jwt_token(session_raw)

    async def test_membership_denials(self):
        """Reject unknown accounts, wrong teams, and nonmembers."""
        for email, team, server in [
            ("missing@example.com", "team-a", "server-a"),
            ("member@example.com", "team-b", "server-a"),
            ("other@example.com", "team-a", "server-a"),
            ("member@example.com", "team-a", "missing"),
        ]:
            with self.subTest(email=email, team=team, server=server), self.assertRaises(HTTPException):
                await member_server(self.db, email, team, server)

    async def test_disabled_entities_and_private_visibility(self):
        """Reject disabled identities, memberships, teams, servers, and private resources."""
        objects = [
            (self.db.query(EmailUser).filter_by(email="member@example.com").one(), "is_active"),
            (self.db.query(EmailTeamMember).one(), "is_active"),
            (self.db.get(EmailTeam, "team-a"), "is_active"),
            (self.db.get(Server, "server-a"), "enabled"),
        ]
        for obj, attribute in objects:
            setattr(obj, attribute, False)
            self.db.commit()
            with self.assertRaises(HTTPException):
                await member_server(self.db, "member@example.com", "team-a", "server-a")
            setattr(obj, attribute, True)
            self.db.commit()
        server = self.db.get(Server, "server-a")
        server.visibility, server.owner_email = "private", "other@example.com"
        self.db.commit()
        with self.assertRaises(HTTPException):
            await member_server(self.db, "member@example.com", "team-a", "server-a")

    async def test_missing_mcp_permission(self):
        """Reject issuance when membership grants no MCP transport permission."""
        self.db.get(Role, "role-a").permissions = ["tools.read"]
        self.db.commit()
        with self.assertRaises(HTTPException):
            await issue_member_key(self.db, "member@example.com", "team-a", "server-a", 7, "self-service")

    async def test_public_server_uses_recipient_team_scope(self):
        """Permit public servers across teams without expanding token team scope."""
        server = self.db.get(Server, "server-a")
        server.team_id, server.visibility = "team-b", "public"
        self.db.commit()
        with patch("mcpgateway.services.server_member_service.get_audit_trail_service"):
            result = await issue_member_key(self.db, "member@example.com", "team-a", "server-a", 7, "self-service")
        claims = jwt.decode(result["api_key"], options={"verify_signature": False})
        self.assertEqual(claims["teams"], ["team-a"])
        server.visibility = "team"
        self.db.commit()
        with self.assertRaises(HTTPException):
            await member_server(self.db, "member@example.com", "team-a", "server-a")

    async def test_password_options_and_issue_boundaries(self):
        """Require real passwords on both endpoints and honor account lockout."""
        # Standard
        from datetime import timedelta

        # Third-Party
        from argon2 import PasswordHasher
        from fastapi import FastAPI
        import httpx

        # First-Party
        from mcpgateway.db import utc_now
        from mcpgateway.routers.server_members import get_db, router

        user = self.db.query(EmailUser).filter_by(email="member@example.com").one()
        user.password_hash = PasswordHasher().hash("Fixture-password-927!")
        self.db.commit()
        app = FastAPI()
        app.include_router(router)
        app.dependency_overrides[get_db] = lambda: self.db
        csrf = "test-csrf-value-" * 3
        with patch("mcpgateway.routers.server_members.settings.self_service_api_keys_enabled", True), patch("mcpgateway.routers.server_members.settings.email_auth_enabled", True):
            async with httpx.AsyncClient(
                transport=httpx.ASGITransport(app=app), base_url="http://test", headers={"Origin": "http://test", "X-CSRF-Token": csrf}, cookies={"mcpgateway_csrf_token": csrf}
            ) as client:
                body = {"email": user.email, "team_id": "team-a", "server_id": "server-a"}
                for endpoint in ("options", "issue"):
                    missing = await client.post(f"/api-key/{endpoint}", json=body)
                    self.assertEqual(missing.status_code, 422)
                    wrong = await client.post(f"/api-key/{endpoint}", json={**body, "password": "wrong-password"})
                    self.assertEqual(wrong.status_code, 403)
                self.db.refresh(user)
                self.assertEqual(user.failed_login_attempts, 2)
                valid = {**body, "password": "Fixture-password-927!"}
                options = await client.post("/api-key/options", json=valid)
                self.assertEqual(options.status_code, 200, options.text)
                self.assertEqual(options.json()["servers"][0]["eligible_team_ids"], ["team-a"])
                user.locked_until = utc_now() + timedelta(minutes=5)
                self.db.commit()
                for endpoint in ("options", "issue"):
                    locked = await client.post(f"/api-key/{endpoint}", json=valid)
                    self.assertEqual(locked.status_code, 403)

    async def test_sso_member_key_boundaries(self):
        """Accept only a short-lived Key proof and enforce CSRF and team access."""
        # Third-Party
        from fastapi import FastAPI
        import httpx

        # First-Party
        from mcpgateway.routers.server_members import get_db, router
        from mcpgateway.utils.member_key_proof import create_member_key_proof

        user = self.db.query(EmailUser).filter_by(email="member@example.com").one()
        user.auth_provider = "google"
        self.db.commit()
        app = FastAPI()
        app.include_router(router)
        app.dependency_overrides[get_db] = lambda: self.db
        csrf = "test-csrf-value-" * 3
        body = {"team_id": "team-a", "server_id": "server-a", "days": 7}
        with patch("mcpgateway.routers.server_members.settings.self_service_api_keys_enabled", True), patch("mcpgateway.routers.server_members.settings.email_auth_enabled", True), patch("mcpgateway.routers.server_members.settings.sso_enabled", True):
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app, client=("127.0.0.2", 12345)), base_url="http://test", headers={"Origin": "http://test", "X-CSRF-Token": csrf}, cookies={"mcpgateway_csrf_token": csrf}) as client:
                for endpoint in ("options", "issue"):
                    missing_cookie = await client.post(f"/api-key/sso/{endpoint}", json=body)
                    self.assertEqual(missing_cookie.status_code, 401)
                client.cookies.set("jwt_token", "fixture-session")
                self.assertEqual((await client.post("/api-key/sso/options", json={})).status_code, 401)
                with patch("mcpgateway.routers.server_members.settings.sso_enabled", False):
                    self.assertEqual((await client.post("/api-key/sso/options", json={})).status_code, 404)
                client.cookies.set("member_key_proof", create_member_key_proof(user.email, "google"))
                self.assertEqual((await client.post("/api-key/sso/options", json={}, headers={"Authorization": "Bearer fixture-session"})).status_code, 403)
                client.cookies.set("member_key_proof", create_member_key_proof(user.email, "other-provider"))
                self.assertEqual((await client.post("/api-key/sso/options", json={})).status_code, 403)
                client.cookies.set("member_key_proof", create_member_key_proof(user.email, "google"))
                no_csrf = await client.post("/api-key/sso/options", json={}, headers={"X-CSRF-Token": "wrong"})
                self.assertEqual(no_csrf.status_code, 403)
                user.auth_provider = "local"
                self.db.commit()
                self.assertEqual((await client.post("/api-key/sso/options", json={})).status_code, 403)
                user.auth_provider = "google"
                self.db.commit()
                options = await client.post("/api-key/sso/options", json={})
                self.assertEqual(options.status_code, 200, options.text)
                self.assertEqual(options.json()["email"], user.email)
                self.assertEqual(options.json()["servers"][0]["eligible_team_ids"], ["team-a"])
                wrong_team = await client.post("/api-key/sso/issue", json={**body, "team_id": "team-b"})
                self.assertEqual(wrong_team.status_code, 403)
                async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app, client=("127.0.0.4", 12345)), base_url="http://test", headers={"Origin": "http://test", "X-CSRF-Token": csrf}, cookies=dict(client.cookies)) as issue_client:
                    with patch("mcpgateway.services.server_member_service.get_audit_trail_service"):
                        issued = await issue_client.post("/api-key/sso/issue", json=body)
                self.assertEqual(issued.status_code, 200, issued.text)
                self.assertEqual(issued.json()["email"], user.email)
                user.is_active = False
                self.db.commit()
                self.assertEqual((await client.post("/api-key/sso/options", json={})).status_code, 403)

    async def test_sso_public_key_without_team(self):
        """Complete the SSO options and issue flow for a public MCP."""
        from fastapi import FastAPI
        import httpx

        from mcpgateway.routers.server_members import get_db, router
        from mcpgateway.utils.member_key_proof import create_member_key_proof

        user = EmailUser(email="public@example.com", auth_provider="google")
        self.db.add(user)
        self.db.add(Role(id="role-public", name="public-viewer-test", scope="global", created_by="member@example.com", permissions=["servers.use", "tools.read"]))
        self.db.flush()
        self.db.add(UserRole(user_email=user.email, role_id="role-public", scope="global", granted_by="member@example.com"))
        self.db.add(Server(id="server-public", name="Public", team_id="team-a", visibility="public", owner_email="member@example.com"))
        self.db.commit()

        app = FastAPI()
        app.include_router(router)
        app.dependency_overrides[get_db] = lambda: self.db
        csrf = "test-csrf-value-" * 3
        with patch("mcpgateway.routers.server_members.settings.self_service_api_keys_enabled", True), patch("mcpgateway.routers.server_members.settings.sso_enabled", True):
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test", headers={"Origin": "http://test", "X-CSRF-Token": csrf}, cookies={"mcpgateway_csrf_token": csrf}) as client:
                self.assertEqual((await client.post("/api-key/sso/options", json={})).status_code, 401)
                client.cookies.set("member_key_proof", create_member_key_proof(user.email, "google"))
                options = await client.post("/api-key/sso/options", json={})
                self.assertEqual(options.status_code, 200, options.text)
                self.assertEqual(options.json()["teams"], [])
                self.assertEqual(options.json()["servers"][0]["id"], "server-public")
                self.assertTrue(options.json()["servers"][0]["public_access"])
                self.assertEqual((await client.post("/api-key/sso/issue", json={"team_id": None, "server_id": "server-a"})).status_code, 403)
                with patch("mcpgateway.services.server_member_service.get_audit_trail_service"):
                    issued = await client.post("/api-key/sso/issue", json={"team_id": None, "server_id": "server-public", "days": 7})
                self.assertEqual(issued.status_code, 200, issued.text)
                self.assertIsNone(issued.json()["team_id"])

    async def test_sso_key_flow_does_not_register_unknown_user(self):
        """Verify existing OAuth accounts without creating or relinking users."""
        # First-Party
        from mcpgateway.services.sso_service import SSOService

        service = SSOService(self.db)
        provider = SimpleNamespace(id="google", is_enabled=True, trusted_domains=["example.com"])
        with patch.object(service, "get_provider", return_value=provider):
            self.assertIsNone(service.authorize_existing_member_key_user({"email": "unknown@example.com", "email_verified": True, "provider": "google"}, "google"))
            self.assertIsNone(self.db.query(EmailUser).filter_by(email="unknown@example.com").one_or_none())
            user = self.db.query(EmailUser).filter_by(email="member@example.com").one()
            user.auth_provider = "google"
            self.db.commit()
            self.assertIs(service.authorize_existing_member_key_user({"email": user.email, "email_verified": True, "provider": "google"}, "google"), user)
            self.assertIsNone(service.authorize_existing_member_key_user({"email": user.email, "email_verified": False, "provider": "google"}, "google"))
            self.assertIsNone(service.authorize_existing_member_key_user({"email": user.email, "email_verified": True, "provider": "other"}, "google"))
            user.auth_provider = "local"
            self.db.commit()
            self.assertIsNone(service.authorize_existing_member_key_user({"email": user.email, "email_verified": True, "provider": "google"}, "google"))

    async def test_sso_member_key_reads_verified_browser_cookie(self):
        """Reject a gateway session while accepting the Key-only OAuth proof."""
        # Third-Party
        from fastapi import FastAPI
        import httpx

        # First-Party
        from mcpgateway.routers.server_members import get_db, router
        from mcpgateway.utils.member_key_proof import create_member_key_proof

        user = self.db.query(EmailUser).filter_by(email="member@example.com").one()
        user.auth_provider = "google"
        self.db.commit()
        app = FastAPI()
        app.include_router(router)
        app.dependency_overrides[get_db] = lambda: self.db
        csrf = "test-csrf-value-" * 3

        with patch("mcpgateway.routers.server_members.settings.self_service_api_keys_enabled", True), patch("mcpgateway.routers.server_members.settings.sso_enabled", True):
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app, client=("127.0.0.3", 12345)), base_url="http://test", headers={"Origin": "http://test", "X-CSRF-Token": csrf, "Accept": "text/html"}, cookies={"mcpgateway_csrf_token": csrf}) as client:
                missing = await client.post("/api-key/sso/options", json={})
                self.assertEqual(missing.status_code, 401)
                client.cookies.set("jwt_token", "fixture-session")
                self.assertEqual((await client.post("/api-key/sso/options", json={})).status_code, 401)
                client.cookies.set("member_key_proof", create_member_key_proof(user.email, "google"))
                options = await client.post("/api-key/sso/options", json={})
                self.assertEqual(options.status_code, 200, options.text)
                self.assertEqual(options.json()["email"], user.email)
                bearer = await client.post("/api-key/sso/options", json={}, headers={"Authorization": "Bearer fixture-session"})
                self.assertEqual(bearer.status_code, 403)

    def test_roster_and_method_filter(self):
        """Include unused members and filter records by their actual server."""
        all_users = member_usage_summary(self.db, 7, None, None, None, None)
        self.assertEqual(all_users["member_states"], {"active": 0, "errors": 0, "idle": 2})
        data = member_usage_summary(self.db, 7, "server-a", None, None, None)
        self.assertEqual(data["members"][0]["calls"], 0)
        self.db.add_all(
            [
                TokenUsageLog(
                    token_jti="deleted-token",
                    user_email="member@example.com",
                    endpoint="/mcp/",
                    status_code=200,
                    response_time_ms=0,
                    mcp_details={"method": "tools/list", "server_id": "server-a", "outcome": "success"},
                ),
                TokenUsageLog(token_jti="deleted-token", user_email="member@example.com", endpoint="/mcp/", status_code=403, mcp_details={"method": "tools/list", "server_id": "server-b"}),
            ]
        )
        self.db.commit()
        data = member_usage_summary(self.db, 7, "server-a", None, None, "tools/list")
        self.assertEqual(data["summary"]["calls"], 1)
        self.assertEqual(data["summary"]["errors"], 0)
        self.assertEqual(data["recent"][0]["server"], "A")


class TestMCPCapture(unittest.TestCase):
    """Verify protocol errors, reported usage, and capture bounds."""

    def test_sse_tool_error_with_usage(self):
        """Read a chunked SSE result without storing tool arguments."""
        capture = MCPUsageCapture()
        capture.append(json.dumps({"jsonrpc": "2.0", "id": 1, "method": "tools/call", "params": {"name": "search", "arguments": {"secret": "not-stored"}}}).encode())
        response = b'data: {"jsonrpc":"2.0","id":1,"result":{"isError":true,"_meta":{"usage":{"input_tokens":12,"output_tokens":4}}}}\n\n'
        capture.append(response[:40], response=True)
        capture.append(response[40:], response=True)
        details = capture.details()
        self.assertEqual(details["outcome"], "error")
        self.assertEqual(details["input_tokens"], 12)
        self.assertNotIn("not-stored", str(details))

    def test_missing_usage_and_bound(self):
        """Keep missing model usage distinct from zero and bound memory."""
        capture = MCPUsageCapture()
        capture.append(b'{"jsonrpc":"2.0","id":1,"method":"tools/list"}')
        capture.append(b'{"jsonrpc":"2.0","id":1,"result":{"tools":[]}}', response=True)
        self.assertNotIn("input_tokens", capture.details())
        capture.append(b"x" * 300000, response=True)
        self.assertEqual(len(capture.response), 262144)
        self.assertTrue(capture.truncated)

    def test_metrics_count_protocol_errors_and_zero_latency(self):
        """Count HTTP-200 MCP errors and preserve zero-millisecond latency."""
        rows = [
            SimpleNamespace(response_time_ms=0, status_code=200, mcp_details={"outcome": "error"}, blocked=False, timestamp=1),
            SimpleNamespace(response_time_ms=100, status_code=200, mcp_details={"input_tokens": 0, "output_tokens": 3}, blocked=False, timestamp=2),
        ]
        data = summarize(rows)
        self.assertEqual(data["avg_ms"], 50)
        self.assertEqual(data["p95_ms"], 100)
        self.assertEqual(data["success_rate"], 50)
        self.assertEqual(data["input_tokens"], 0)
        self.assertEqual(data["usage_reported_calls"], 1)
        self.assertIsNone(summarize(rows[:1])["input_tokens"])


class TestManagementBoundaries(unittest.TestCase):
    """Check feature flags and management-plane isolation."""

    def test_disabled_self_service(self):
        """Reject public issuance when either feature flag is disabled."""
        # First-Party
        from mcpgateway.routers.server_members import require_self_service

        for enabled, email_auth in ((False, True), (True, False)):
            with patch("mcpgateway.routers.server_members.settings") as settings:
                settings.self_service_api_keys_enabled = enabled
                settings.email_auth_enabled = email_auth
                settings.sso_enabled = False
                settings.portal_sso_enabled = False
                with self.assertRaises(HTTPException) as error:
                    require_self_service()
                self.assertEqual(error.exception.status_code, 404)

    def test_management_denials(self):
        """Reject anonymous users, API tokens, members, and narrowed admins."""
        # First-Party
        from mcpgateway.routers.server_members import require_platform_session

        for user in (
            {},
            {"auth_method": "anonymous"},
            {"auth_method": "api_token", "is_admin": True},
            {"auth_method": "session", "is_admin": False},
            {"auth_method": "session", "is_admin": True, "token_teams": ["other-team"]},
        ):
            with self.subTest(user=user), self.assertRaises(HTTPException):
                require_platform_session(user)

    def test_metadata_migration(self):
        """Upgrade existing data twice and downgrade without changing request fields."""
        # Third-Party
        from alembic.migration import MigrationContext
        from alembic.operations import Operations

        migration = importlib.import_module("mcpgateway.alembic.versions.c31d7e4a092b_add_member_mcp_usage_details")
        engine = create_engine("sqlite:///:memory:")
        with engine.begin() as connection:
            connection.execute(text("CREATE TABLE token_usage_logs (id INTEGER PRIMARY KEY, user_email VARCHAR(255))"))
            connection.execute(text("INSERT INTO token_usage_logs VALUES (1, 'member@example.com')"))
            with Operations.context(MigrationContext.configure(connection)):
                migration.upgrade()
                migration.upgrade()
                self.assertIn("mcp_details", {c["name"] for c in inspect(connection).get_columns("token_usage_logs")})
                migration.downgrade()
                migration.downgrade()
            self.assertEqual(connection.scalar(text("SELECT user_email FROM token_usage_logs WHERE id=1")), "member@example.com")
        engine.dispose()

    def test_encrypted_token_migration(self):
        """Preserve historical hashes through repeated upgrades and downgrades."""
        from alembic.migration import MigrationContext
        from alembic.operations import Operations

        migration = importlib.import_module("mcpgateway.alembic.versions.d72f8a1c903e_store_encrypted_api_tokens")
        engine = create_engine("sqlite:///:memory:")
        with engine.begin() as connection:
            connection.execute(text("CREATE TABLE email_api_tokens (id INTEGER PRIMARY KEY, token_hash VARCHAR(255))"))
            connection.execute(text("INSERT INTO email_api_tokens VALUES (1, 'historical-hash')"))
            with Operations.context(MigrationContext.configure(connection)):
                migration.upgrade()
                migration.upgrade()
                self.assertIsNone(connection.scalar(text("SELECT encrypted_token FROM email_api_tokens WHERE id=1")))
                migration.downgrade()
                migration.downgrade()
            self.assertEqual(connection.scalar(text("SELECT token_hash FROM email_api_tokens WHERE id=1")), "historical-hash")
        engine.dispose()


class TestUsageMiddleware(unittest.IsolatedAsyncioTestCase):
    """Verify that collection preserves ASGI messages and scoped server identity."""

    async def test_stream_passthrough(self):
        """Capture a protocol error without changing response bytes."""
        # First-Party
        from mcpgateway.middleware.token_usage_middleware import TokenUsageMiddleware

        payload = b'{"jsonrpc":"2.0","id":1,"method":"tools/call","params":{"name":"example"}}'
        output = b'data: {"jsonrpc":"2.0","id":1,"result":{"isError":true}}\n\n'
        received = []

        async def app(scope, receive, send):
            self.assertEqual((await receive())["body"], payload)
            await send({"type": "http.response.start", "status": 200, "headers": []})
            await send({"type": "http.response.body", "body": output[:20], "more_body": True})
            await send({"type": "http.response.body", "body": output[20:], "more_body": False})

        async def receive():
            return {"type": "http.request", "body": payload}

        async def send(message):
            received.append(message)

        scope = {
            "type": "http",
            "method": "POST",
            "path": "/mcp/",
            "modified_path": "/servers/server-a/mcp",
            "headers": [],
            "state": {"auth_method": "api_token", "jti": "test-jti", "user_email": "member@example.com"},
        }
        with patch("mcpgateway.middleware.token_usage_middleware.fresh_db_session"), patch("mcpgateway.middleware.token_usage_middleware.TokenCatalogService") as service:
            service.return_value.log_token_usage = AsyncMock()
            await TokenUsageMiddleware(app)(scope, receive, send)
            details = service.return_value.log_token_usage.call_args.kwargs["mcp_details"]
        self.assertEqual(details["server_id"], "server-a")
        self.assertEqual(details["outcome"], "error")
        self.assertEqual(b"".join(m.get("body", b"") for m in received), output)


if __name__ == "__main__":
    unittest.main()
