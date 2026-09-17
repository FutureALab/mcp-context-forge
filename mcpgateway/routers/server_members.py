"""Serve member imports, key issuance, and the member usage panel.

Copyright contributors to the MCP-CONTEXT-FORGE project.
SPDX-License-Identifier: Apache-2.0
"""

# Standard
from datetime import timedelta
import io
import secrets

# Third-Party
from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import JSONResponse, RedirectResponse, Response
import openpyxl
from pydantic import BaseModel, EmailStr, Field, SecretStr
from sqlalchemy import func, select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

# First-Party
from mcpgateway.admin import (
    _request_origin_matches,
    _set_admin_csrf_cookie,
    ADMIN_CSRF_COOKIE_NAME,
    ADMIN_CSRF_HEADER_NAME,
    get_bundle_css_files,
    get_db,
    parse_bulk_user_workbook,
    rate_limit,
)
from mcpgateway.auth_context import get_user_email
from mcpgateway.config import settings
from mcpgateway.db import EmailApiToken, EmailTeam, EmailTeamMember, EmailUser, Server, utc_now
from mcpgateway.middleware.rbac import get_current_user_with_permissions, require_permission
from mcpgateway.routers.tokens import _require_authenticated_session
from mcpgateway.services.email_auth_service import EmailAuthService
from mcpgateway.services.server_member_service import issue_member_key, registered_user
from mcpgateway.services.server_service import ServerService
from mcpgateway.services.team_management_service import MemberAlreadyExistsError, TeamManagementError, TeamManagementService

router = APIRouter()


class AccountRequest(BaseModel):
    """An existing account supplied by the visitor."""

    email: EmailStr
    password: SecretStr = Field(min_length=1, max_length=1024)


class MemberKeyRequest(AccountRequest):
    """A server-scoped key request with a bounded lifetime."""

    team_id: str = Field(min_length=1, max_length=36)
    server_id: str = Field(min_length=1, max_length=36)
    days: int = Field(default=30, ge=1, le=365)


class BulkKeyRequest(BaseModel):
    """Select one server or all team-backed servers."""

    server_id: str | None = Field(default=None, max_length=36)
    team_id: str | None = Field(default=None, max_length=36)
    days: int = Field(default=30, ge=1, le=365)


class KeyExportRequest(BaseModel):
    """Hold the current browser's bounded key issuance results."""

    results: list[dict] = Field(min_length=1, max_length=500)


def management_team(db: Session, server: Server, team_id: str | None) -> EmailTeam:
    """Resolve a shared team without changing server ownership or visibility."""
    team = db.get(EmailTeam, team_id or server.team_id) if team_id or server.team_id else None
    if not team or not team.is_active or team.is_personal:
        raise HTTPException(400, "请选择启用的共享团队；个人空间不能添加其他成员")
    if server.visibility != "public" and server.team_id != team.id:
        raise HTTPException(403, "非公众 MCP 只能管理其所属团队，请先在服务器设置中调整归属和可见性")
    return team


def require_self_service() -> None:
    """Reject self-service when the operator disables the feature."""
    if not settings.self_service_api_keys_enabled or not settings.email_auth_enabled:
        raise HTTPException(404, "自助领取未启用")


def require_platform_session(user: dict) -> str:
    """Restrict member administration to unrestricted platform sessions."""
    _require_authenticated_session(user)
    if not user.get("is_admin") or user.get("token_teams") is not None:
        raise HTTPException(403, "需要平台管理员权限")
    return get_user_email(user)


def require_public_csrf(request: Request) -> None:
    """Check browser origin and the pre-authentication CSRF cookie."""
    cookie = request.cookies.get(ADMIN_CSRF_COOKIE_NAME, "")
    header = request.headers.get(ADMIN_CSRF_HEADER_NAME, "")
    if not _request_origin_matches(request) or len(cookie) < 32 or not secrets.compare_digest(cookie, header):
        raise HTTPException(403, "请刷新页面后重试")


async def authenticate_member(request: Request, body: AccountRequest, db: Session) -> EmailUser:
    """Verify the local password with the existing account lockout policy."""
    user = await EmailAuthService(db).authenticate_user(
        str(body.email), body.password.get_secret_value(), ip_address=request.client.host if request.client else None, user_agent=request.headers.get("user-agent")
    )
    if user is None:
        raise HTTPException(403, "邮箱或密码错误，或账号不可用")
    return user


@router.get("/api-key")
async def member_key_page(request: Request):
    """Render the public key page using the login template."""
    require_self_service()
    response = request.app.state.templates.TemplateResponse(
        request,
        "member_key.html",
        {
            "root_path": request.scope.get("root_path", ""),
            "bundle_css": get_bundle_css_files(),
        },
    )
    _set_admin_csrf_cookie(request, response)
    response.headers["Cache-Control"] = "no-store"
    return response


@router.post("/api-key/options")
@rate_limit(requests_per_minute=20)
async def member_key_options(request: Request, body: AccountRequest, db: Session = Depends(get_db)):
    """List active teams and servers accessible to the supplied account."""
    require_self_service()
    require_public_csrf(request)
    user = await authenticate_member(request, body, db)
    teams = (
        db.execute(
            select(EmailTeam)
            .join(EmailTeamMember)
            .where(
                EmailTeamMember.user_email == user.email,
                EmailTeamMember.is_active.is_(True),
                EmailTeam.is_active.is_(True),
            )
            .order_by(EmailTeam.name)
        )
        .scalars()
        .all()
    )
    servers = (
        db.execute(
            select(Server).where(Server.enabled.is_(True), (Server.team_id.in_([t.id for t in teams])) | (Server.visibility == "public") | (Server.owner_email == user.email)).order_by(Server.name)
        )
        .scalars()
        .all()
    )
    visible = []
    service = ServerService()
    for server in servers:
        eligible = [team.id for team in teams if await service._check_server_access(db, server, user.email, [team.id])]
        if eligible:
            visible.append({"id": server.id, "name": server.name, "team_id": server.team_id, "visibility": server.visibility, "eligible_team_ids": eligible})
    return JSONResponse({"teams": [{"id": t.id, "name": t.name} for t in teams], "servers": visible}, headers={"Cache-Control": "no-store"})


@router.post("/api-key/issue")
@rate_limit(requests_per_minute=10)
async def member_key_issue(request: Request, body: MemberKeyRequest, db: Session = Depends(get_db)):
    """Issue a key after password, membership, and visibility checks."""
    require_self_service()
    require_public_csrf(request)
    user = await authenticate_member(request, body, db)
    recent = db.scalar(
        select(func.count())
        .select_from(EmailApiToken)
        .where(
            EmailApiToken.user_email == user.email,
            EmailApiToken.created_at >= utc_now() - timedelta(hours=1),
        )
    )
    if recent >= 10:
        raise HTTPException(429, "该账号每小时最多生成 10 个 Key，请稍后重试")
    result = await issue_member_key(db, user.email, body.team_id, body.server_id, body.days, "self-service")
    return JSONResponse(result, headers={"Cache-Control": "no-store", "Pragma": "no-cache"})


@router.get("/member-usage")
@require_permission("admin.overview")
async def member_usage_page(request: Request, user=Depends(get_current_user_with_permissions)):
    """Open member analytics inside the existing administration interface."""
    require_platform_session(user)
    return RedirectResponse(f"{request.scope.get('root_path', '')}/admin/#member-usage", status_code=303)


@router.get("/member-usage/servers")
@require_permission("servers.read")
async def member_servers(user=Depends(get_current_user_with_permissions), db: Session = Depends(get_db)):
    """List servers with their teams for the member management panel."""
    require_platform_session(user)
    rows = db.execute(select(Server, EmailTeam).outerjoin(EmailTeam, Server.team_id == EmailTeam.id).order_by(Server.name)).all()
    return {
        "teams": [{"id": t.id, "name": t.name} for t in db.execute(select(EmailTeam).where(EmailTeam.is_active.is_(True), EmailTeam.is_personal.is_(False)).order_by(EmailTeam.name)).scalars()],
        "servers": [
            {
                "id": s.id,
                "name": s.name,
                "team_id": s.team_id,
                "team_name": t.name if t else None,
                "visibility": s.visibility,
                "is_personal": bool(t and t.is_personal),
                "enabled": s.enabled,
                "can_import": bool(t and t.is_active and not t.is_personal),
            }
            for s, t in rows
        ],
    }


@router.get("/member-usage/template")
@require_permission("teams.manage_members")
async def member_import_template(user=Depends(get_current_user_with_permissions)):
    """Download the workbook headers for existing member accounts."""
    require_platform_session(user)
    workbook = openpyxl.Workbook()
    workbook.active.append(["邮箱"])
    content = io.BytesIO()
    workbook.save(content)
    workbook.close()
    return Response(content.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="members.xlsx"'})


@router.post("/member-usage/import/{server_id}")
@require_permission("teams.manage_members")
async def member_import(request: Request, server_id: str, user=Depends(get_current_user_with_permissions), db: Session = Depends(get_db)):
    """Add registered spreadsheet accounts to the selected server's team."""
    actor = require_platform_session(user)
    server = db.get(Server, server_id)
    if server is None or not server.enabled:
        raise HTTPException(400, "请选择启用的虚拟 MCP")
    form = await request.form()
    team = management_team(db, server, str(form.get("team_id") or "") or None)
    upload = form.get("file")
    if not hasattr(upload, "read") or not str(upload.filename).lower().endswith(".xlsx"):
        raise HTTPException(400, "请选择 .xlsx 文件")
    content = await upload.read(5 * 1024 * 1024 + 1)
    await upload.close()
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(400, "Excel 文件不能超过 5 MB")
    try:
        rows = parse_bulk_user_workbook(content)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    results = []
    for index, row in enumerate(rows, 2):
        email = str(row.get("email") or "").strip()
        try:
            target = registered_user(db, email)
            await TeamManagementService(db).add_member_to_team(team.id, target.email, role="member", invited_by=actor, grant_source="manual")
            status, message = "added", "已加入 Team"
        except MemberAlreadyExistsError:
            status, message = "skipped", "已是 Team 成员"
        except HTTPException:
            status, message = "failed", "账号未注册或已停用，请先在用户管理中创建账号"
        except TeamManagementError as exc:
            db.rollback()
            status, message = "failed", str(exc)
        results.append({"row": index, "email": email, "status": status, "message": message})
    return {"results": results}


@router.post("/member-usage/keys")
@require_permission("tokens.create")
async def member_bulk_keys(body: BulkKeyRequest, user=Depends(get_current_user_with_permissions), db: Session = Depends(get_db)):
    """Generate a different key for each active server and member pair."""
    actor = require_platform_session(user)
    query = select(Server).where(Server.enabled.is_(True))
    if body.server_id:
        query = query.where(Server.id == body.server_id)
    servers = db.execute(query).scalars().all()
    if body.server_id and not servers:
        raise HTTPException(404, "虚拟 MCP 不存在、已停用或未关联 Team")
    pairs = []
    for server in servers:
        if not body.server_id and body.team_id is None:
            owning_team = db.get(EmailTeam, server.team_id) if server.team_id else None
            if not owning_team or owning_team.is_personal or not owning_team.is_active:
                continue
        team = management_team(db, server, body.team_id)
        emails = (
            db.execute(
                select(EmailTeamMember.user_email)
                .join(EmailUser, EmailTeamMember.user_email == EmailUser.email)
                .where(
                    EmailTeamMember.team_id == team.id,
                    EmailTeamMember.is_active.is_(True),
                    EmailUser.is_active.is_(True),
                )
            )
            .scalars()
            .all()
        )
        pairs.extend((server, email, team.id) for email in emails)
    if len(pairs) > 500:
        raise HTTPException(400, "单次最多生成 500 个 Key，请按虚拟 MCP 分批生成")
    results = []
    for server, email, team_id in pairs:
        try:
            result = await issue_member_key(db, email, team_id, server.id, body.days, actor)
            results.append({"status": "renewed" if result.get("renewed") else "created", **result})
        except (HTTPException, ValueError) as exc:
            db.rollback()
            results.append({"status": "failed", "email": email, "server_id": server.id, "message": exc.detail if isinstance(exc, HTTPException) else str(exc)})
        except SQLAlchemyError:
            db.rollback()
            results.append({"status": "failed", "email": email, "server_id": server.id, "message": "数据库写入失败，请稍后重试此成员"})
    return JSONResponse({"results": results}, headers={"Cache-Control": "no-store", "Pragma": "no-cache"})


@router.post("/member-usage/keys/export")
@require_permission("tokens.create")
async def member_keys_export(body: KeyExportRequest, user=Depends(get_current_user_with_permissions)):
    """Export transient issuance results as text cells without persisting raw keys."""
    require_platform_session(user)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "成员 API Key"
    columns = [
        ("成员邮箱", "email"),
        ("虚拟 MCP", "server_name"),
        ("服务器 ID", "server_id"),
        ("团队 ID", "team_id"),
        ("状态", "status"),
        ("API Key", "api_key"),
        ("到期时间", "expires_at"),
        ("说明", "message"),
    ]
    sheet.append([label for label, _ in columns])
    for row in body.results:
        values = [str(row.get(key) or "") for _, key in columns]
        if any(len(value) > 16384 for value in values):
            raise HTTPException(400, "导出字段过长，请分批操作")
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in "ABCDEFGH":
        sheet.column_dimensions[column].width = 32 if column != "F" else 60
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return Response(
        output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="mcp-member-keys.xlsx"', "Cache-Control": "no-store", "Pragma": "no-cache"},
    )


@router.get("/member-usage/data")
@require_permission("tokens.read")
async def member_usage_data(
    days: int = Query(7, ge=1, le=90),
    server_id: str | None = None,
    email: str | None = None,
    token_id: str | None = None,
    method: str | None = None,
    user=Depends(get_current_user_with_permissions),
    db: Session = Depends(get_db),
):
    """Return filtered member and method metrics from recorded requests."""
    require_platform_session(user)
    # First-Party
    from mcpgateway.services.member_usage_service import member_usage_summary

    return member_usage_summary(db, days, server_id, email, token_id, method)
